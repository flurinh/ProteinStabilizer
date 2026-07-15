#!/usr/bin/env python3
"""Normalize BenchmarkStabilityDatasets.xlsx into a long fine-tuning table."""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "BenchmarkStabilityDatasets.xlsx"
OUTPUT = ROOT / "data" / "curated" / "gpcr_finetune.csv"

FIELDS = [
    "protein_id",
    "uniprot_id",
    "species",
    "sequence",
    "is_membrane",
    "assay_id",
    "mutation",
    "raw_mutation",
    "mutation_note",
    "wt_aa",
    "position",
    "mut_aa",
    "is_wt",
    "stability_percent",
    "stability_delta_percent",
    "source_url",
]


def fasta_sequence(value: Any) -> str:
    return "".join(
        line.strip()
        for line in str(value).splitlines()
        if line.strip() and not line.startswith(">")
    )


def parse_mutation(raw: str, infer_alanine_scan: bool = False) -> dict[str, Any]:
    if raw.lower() == "wt":
        return {
            "mutation": "WT",
            "raw_mutation": raw,
            "mutation_note": "",
            "wt_aa": "",
            "position": "",
            "mut_aa": "",
            "is_wt": 1,
        }

    match = re.fullmatch(r"([A-Z])(\d+)([A-Z])?([^A-Z0-9]*)", raw)
    if not match:
        raise ValueError(f"unrecognized mutation: {raw!r}")
    wt_aa, position_text, mut_aa, note = match.groups()
    if mut_aa is None:
        if not infer_alanine_scan:
            raise ValueError(f"mutation destination is missing: {raw!r}")
        mut_aa = "L" if wt_aa == "A" else "A"
    position = int(position_text)
    return {
        "mutation": f"{wt_aa}{position}{mut_aa}",
        "raw_mutation": raw,
        "mutation_note": note,
        "wt_aa": wt_aa,
        "position": position,
        "mut_aa": mut_aa,
        "is_wt": 0,
    }


def add_row(
    rows: list[dict[str, Any]],
    *,
    protein_id: str,
    uniprot_id: str,
    species: str,
    sequence: str,
    assay_id: str,
    raw_mutation: str,
    stability_percent: float,
    source_url: str,
    infer_alanine_scan: bool = False,
) -> None:
    mutation = parse_mutation(raw_mutation, infer_alanine_scan)
    if not mutation["is_wt"]:
        position = int(mutation["position"])
        observed = sequence[position - 1]
        if observed != mutation["wt_aa"]:
            raise ValueError(
                f"{protein_id} {raw_mutation}: sequence has {observed} at {position}"
            )
    rows.append(
        {
            "protein_id": protein_id,
            "uniprot_id": uniprot_id,
            "species": species,
            "sequence": sequence,
            "is_membrane": 1,
            "assay_id": assay_id,
            **mutation,
            "stability_percent": stability_percent,
            "stability_delta_percent": stability_percent - 50.0,
            "source_url": source_url,
        }
    )


def main() -> None:
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []

    ws = workbook["neurotensin"]
    sequence = fasta_sequence(ws["A2"].value)
    source = str(ws["A3"].value)
    for row in range(6, 41):
        raw = str(ws.cell(row, 1).value)
        for assay_id, column in (("ntr_apo", 2), ("ntr_agonist_bound", 3)):
            add_row(
                rows,
                protein_id="NTSR1_RAT",
                uniprot_id="P20789",
                species="Rattus norvegicus",
                sequence=sequence,
                assay_id=assay_id,
                raw_mutation=raw,
                stability_percent=float(ws.cell(row, column).value),
                source_url=source,
            )

    ws = workbook["A2A"]
    sequence = fasta_sequence(ws["A2"].value)
    source = str(ws["A3"].value)
    for assay_id, start, end in (
        ("a2a_agonist_readout", 7, 34),
        ("a2a_antagonist_readout", 38, 55),
    ):
        for row in range(start, end + 1):
            add_row(
                rows,
                protein_id="AA2AR_HUMAN",
                uniprot_id="P29274",
                species="Homo sapiens",
                sequence=sequence,
                assay_id=assay_id,
                raw_mutation=str(ws.cell(row, 1).value),
                stability_percent=float(ws.cell(row, 2).value),
                source_url=source,
            )

    ws = workbook["β1AR"]
    sequence = fasta_sequence(ws["B26"].value)
    source = "https://doi.org/10.1073/pnas.0711253105"
    for row in (*range(31, 43), *range(44, 53)):
        add_row(
            rows,
            protein_id="ADRB1_MELGA",
            uniprot_id="P07700",
            species="Meleagris gallopavo",
            sequence=sequence,
            assay_id="b1ar_apo_antagonist_readout",
            raw_mutation=str(ws.cell(row, 1).value),
            stability_percent=float(ws.cell(row, 2).value),
            source_url=source,
            infer_alanine_scan=True,
        )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    print(f"wrote {len(rows)} rows to {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
